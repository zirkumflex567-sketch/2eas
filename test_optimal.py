
import openpyxl, datetime

wb = openpyxl.load_workbook(r'C:\Users\zirku\Downloads\check06.xlsx', data_only=True)
ws = wb.active

deals = []
for r in range(100, ws.max_row + 1):
    if ws.cell(r, 5).value == 'out' and ws.cell(r, 4).value in ['buy', 'sell']:
        t_str = ws.cell(r, 1).value
        p = float(ws.cell(r, 11).value or 0.0)
        c = float(ws.cell(r, 9).value or 0.0)
        s = float(ws.cell(r, 10).value or 0.0)
        dt = datetime.datetime.strptime(str(t_str).strip(), '%Y.%m.%d %H:%M:%S')
        deals.append({'date': dt.date(), 'hour': dt.hour, 'net': p + c + s})

# Filter out 15-17
filtered = [d for d in deals if not (15 <= d['hour'] <= 17)]
print(f'Filtered Deals: {len(filtered)}')
print(f'Net Profit: ')
w = sum(d['net'] for d in filtered if d['net'] > 0)
l = sum(d['net'] for d in filtered if d['net'] <= 0)
print(f'Gross Win: ')
print(f'Gross Loss: ')
print(f'Profit Factor: {abs(w/l):.2f}')

# Daily DD check
days = {}
cur_bal = 10000.0
peak_bal = 10000.0
max_total_dd = 0.0
worst_days = []

for d in filtered:
    day = d['date']
    if day not in days:
        days[day] = {'start': cur_bal, 'pnl': 0.0, 'min': cur_bal}
    cur_bal += d['net']
    days[day]['pnl'] += d['net']
    if cur_bal < days[day]['min']:
        days[day]['min'] = cur_bal
    if cur_bal > peak_bal:
        peak_bal = cur_bal
    dd = (peak_bal - cur_bal) / peak_bal * 100.0
    if dd > max_total_dd:
        max_total_dd = dd

for day, inf in days.items():
    s = inf['start']
    m = inf['min']
    dip = (s - m) / s * 100.0
    worst_days.append((day, s, inf['pnl'], dip))

worst_days.sort(key=lambda x: x[3], reverse=True)
print(f'Max Trailing DD: {max_total_dd:.2f}%')
breaches = [x for x in worst_days if x[3] >= 3.5]
print(f'Daily DD breaches (>=3.5%): {len(breaches)}')
for b in breaches:
    print(f'  {b[0]}: Start=, PnL=, Dip={b[3]:.2f}%')
