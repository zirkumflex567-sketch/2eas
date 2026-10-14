import openpyxl, datetime
import pandas as pd
import numpy as np

# Let's verify what happens when Titan Pro's max loss is capped at $50-$60 instead of $150
p_titan = r'C:\Users\zirku\Downloads\mt5_strategy_research\rev\500er\500acc ohne schutz.xlsx'
wb = openpyxl.load_workbook(p_titan, data_only=True)
ws = wb.active

trades_r = 1517
trades = []
for r in range(trades_r+1, ws.max_row+1):
    c1 = ws.cell(row=r, column=1).value
    if not c1 or not str(c1).startswith('2026'): continue
    direction = ws.cell(row=r, column=5).value
    profit = float(ws.cell(row=r, column=11).value or 0)
    comm = float(ws.cell(row=r, column=9).value or 0)
    swap = float(ws.cell(row=r, column=10).value or 0)
    comment = str(ws.cell(row=r, column=13).value or '')
    if direction == 'out':
        trades.append({
            'time': str(c1),
            'profit': profit + comm + swap,
            'comment': comment
        })

print(f'Total trades analyzed: {len(trades)}')
orig_net = sum(t['profit'] for t in trades)
orig_wins = sum(t['profit'] for t in trades if t['profit'] > 0)
orig_losses = sum(t['profit'] for t in trades if t['profit'] <= 0)

print('=== ORIGINAL TITAN PRO ON 500 USD ===')
print(f'Gross Profit:  ${orig_wins:,.2f}')
print(f'Gross Loss:    ${orig_losses:,.2f}')
print(f'Net Profit:    ${orig_net:,.2f}')
print(f'Profit Factor: {abs(orig_wins/orig_losses):.2f}')

# What if losses are capped at max -$60 (or -$50) like Market EA?
for cap in [-75.0, -60.0, -50.0]:
    capped_losses = sum(max(cap, t['profit']) for t in trades if t['profit'] <= 0)
    sim_net = orig_wins + capped_losses
    sim_pf = abs(orig_wins / capped_losses)
    print(f'\nWith Loss Capped at ${abs(cap):.0f}:')
    print(f'  Gross Loss:    ${capped_losses:,.2f} (Saved: ${abs(orig_losses - capped_losses):,.2f})')
    print(f'  Net Profit:    ${sim_net:,.2f} (vs ${orig_net:,.2f} original)')
    print(f'  Profit Factor: {sim_pf:.2f}')
