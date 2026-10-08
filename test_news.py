
import openpyxl, datetime

wb = openpyxl.load_workbook(r'C:\Users\zirku\Downloads\check06.xlsx', data_only=True)
ws = wb.active

deals = []
for r in range(100, ws.max_row + 1):
    if ws.cell(r, 5).value == 'out' and ws.cell(r, 4).value in ['buy', 'sell']:
        t_str = ws.cell(r, 1).value
        p = float(ws.cell(r, 11).value or 0.0)
        c = float(ws.cell(r, 9).value or 0.0)
        s = float(ws.cell(r, 10).value or 0.0)
        dt = datetime.datetime.strptime(str(t_str).strip(), '%Y.%m.%d %H:%M:%S')
        deals.append({'hour': dt.hour, 'net': p + c + s})

print('check06 Total Net:', sum(d['net'] for d in deals))
# If we filter 15-17
d_no_news = [d for d in deals if not (15 <= d['hour'] <= 17)]
print('check06 Without 15-17 hours:', sum(d['net'] for d in d_no_news))
