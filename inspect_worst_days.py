import openpyxl
import pandas as pd

file_path = r'C:\Users\zirku\Downloads\ReportTester-53087839.xlsx'
wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb['Sheet1']

# Parse all deals
deals = []
for r in range(6600, ws.max_row + 1):
    time_val = ws.cell(r, 1).value
    if not time_val:
        break
    deals.append({
        'row': r,
        'time': str(time_val),
        'trade_id': ws.cell(r, 2).value,
        'symbol': ws.cell(r, 3).value,
        'type': ws.cell(r, 4).value,
        'direction': str(ws.cell(r, 5).value),
        'volume': float(ws.cell(r, 6).value or 0),
        'price': float(ws.cell(r, 7).value or 0),
        'order_id': ws.cell(r, 8).value,
        'commission': float(ws.cell(r, 9).value or 0),
        'swap': float(ws.cell(r, 10).value or 0),
        'profit': float(ws.cell(r, 11).value or 0),
        'balance': float(ws.cell(r, 12).value) if ws.cell(r, 12).value is not None else None,
        'comment': str(ws.cell(r, 13).value or '')
    })

df = pd.DataFrame(deals)
df['datetime'] = pd.to_datetime(df['time'], format='%Y.%m.%d %H:%M:%S', errors='coerce')
df['date'] = df['datetime'].dt.strftime('%Y-%m-%d')

# Let's inspect 2026-09-08
deals_0908 = df[df['date'] == '2026-09-08']
print(f"=== DEALS ON 2026-09-08 (Total: {len(deals_0908)}) ===")
print(deals_0908[['time', 'type', 'direction', 'volume', 'price', 'profit', 'balance', 'comment']].to_string())

# Also inspect comments on 2026-04-23
deals_0423 = df[df['date'] == '2026-04-23']
print(f"\n=== DEALS ON 2026-04-23 (Total: {len(deals_0423)}) ===")
print(deals_0423[['time', 'type', 'direction', 'volume', 'price', 'profit', 'balance', 'comment']].to_string())

# Inspect all comments across the entire backtest to see if 'sl', 'tp', 'close', or risk breach occurred
out_deals = df[df['direction'].str.contains('out', case=False, na=False)]
print("\nUnique comments on exit deals:")
print(out_deals['comment'].value_counts())
