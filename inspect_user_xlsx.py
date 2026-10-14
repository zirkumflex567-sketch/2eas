import openpyxl
import pandas as pd
import numpy as np

file_path = r'C:/Users/zirku/.gemini/antigravity/brain/c98ddc14-d641-4466-a351-59d533463650/.user_uploaded/media_1791964823055_1576354d.xlsx'

wb = openpyxl.load_workbook(file_path, data_only=True)
print(f"Sheet names: {wb.sheetnames}")
ws = wb.active
print(f"Active sheet: {ws.title}, max_row: {ws.max_row}, max_column: {ws.max_column}")

# Search for key sections
for r in range(1, min(100, ws.max_row + 1)):
    val = ws.cell(r, 1).value
    if val:
        val_str = str(val).strip()
        if any(k in val_str for k in ['Ersteinlage', 'Nettogewinn', 'Rückgang', 'Trades', 'Orders']):
            row_vals = [str(ws.cell(r, c).value or '') for c in range(1, 15)]
            print(f"Row {r}: {row_vals[:6]}")
