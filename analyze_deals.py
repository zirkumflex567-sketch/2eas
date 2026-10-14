import openpyxl, os
from collections import defaultdict

folder = r'C:\Users\zirku\Downloads\mt5_strategy_research\rev'
p = os.path.join(folder, '500eurotest.xlsx')
wb = openpyxl.load_workbook(p, data_only=True)
ws = wb.active

# Find trades
trades_r = None
for r in range(1, ws.max_row+1):
    if ws.cell(row=r, column=1).value in ['Trades', 'Deals', 'Geschäfte']:
        trades_r = r
        break

print(f'Trades header at row: {trades_r}')

deals = []
for r in range(trades_r+2, ws.max_row+1):
    c1 = ws.cell(row=r, column=1).value
    if not c1 or not str(c1).startswith('2026'):
        continue
    d = {
        'time': str(c1),
        'ticket': ws.cell(row=r, column=2).value,
        'symbol': ws.cell(row=r, column=3).value,
        'type': ws.cell(row=r, column=4).value,
        'dir': ws.cell(row=r, column=5).value,
        'vol': float(ws.cell(row=r, column=6).value or 0),
        'price': float(ws.cell(row=r, column=7).value or 0),
        'order': ws.cell(row=r, column=8).value,
        'comm': float(ws.cell(row=r, column=9).value or 0),
        'swap': float(ws.cell(row=r, column=10).value or 0),
        'profit': float(ws.cell(row=r, column=11).value or 0),
        'balance': float(ws.cell(row=r, column=12).value or 0),
        'comment': str(ws.cell(row=r, column=13).value or ''),
    }
    deals.append(d)

print(f'Total deals: {len(deals)}')

# Let's inspect first 25 deals
print('\n--- FIRST 25 DEALS ---')
for d in deals[:25]:
    print('%s | %-3s | %-4s | Lot: %.2f | Price: %7.2f | Profit: %7.2f | Bal: %8.2f | %s' % 
          (d['time'], d['dir'], d['type'], d['vol'], d['price'], d['profit'], d['balance'], d['comment']))

# Analyze max open positions
open_pos = 0
max_open = 0
in_count = 0
out_count = 0
sl_count = 0
tp_count = 0
close_count = 0
lots = set()

for d in deals:
    lots.add(d['vol'])
    if d['dir'] == 'in':
        open_pos += 1
        in_count += 1
        if open_pos > max_open:
            max_open = open_pos
    elif d['dir'] == 'out':
        open_pos -= 1
        out_count += 1
        if 'sl' in d['comment'].lower():
            sl_count += 1
        elif 'tp' in d['comment'].lower():
            tp_count += 1
        else:
            close_count += 1

print(f'\nTotal in: {in_count}, Total out: {out_count}')
print(f'Max simultaneous open positions: {max_open}')
print(f'Lots used: {sorted(list(lots))}')
print(f'Exit types: Normal/Signal={close_count}, SL={sl_count}, TP={tp_count}')
