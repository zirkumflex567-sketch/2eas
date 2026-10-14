import openpyxl

p = r'C:\Users\zirku\Downloads\mt5_strategy_research\rev\500eurotest.xlsx'
wb = openpyxl.load_workbook(p, data_only=True)
ws = wb.active

# Orders table: row 62 to 894
orders = []
for r in range(62, 895):
    t_str = ws.cell(row=r, column=1).value
    if not t_str or not str(t_str).startswith('2026'):
        continue
    o_ticket = ws.cell(row=r, column=2).value
    sym = ws.cell(row=r, column=3).value
    typ = ws.cell(row=r, column=4).value
    vol_str = ws.cell(row=r, column=5).value
    sl = ws.cell(row=r, column=8).value
    tp = ws.cell(row=r, column=9).value
    orders.append({
        'ticket': o_ticket,
        'time': str(t_str),
        'type': typ,
        'sl': float(sl) if sl else None,
        'tp': float(tp) if tp else None
    })

# Deals table: row 898 to max_row
deals = []
for r in range(898, ws.max_row+1):
    t_str = ws.cell(row=r, column=1).value
    if not t_str or not str(t_str).startswith('2026'):
        continue
    d_ticket = ws.cell(row=r, column=2).value
    typ = ws.cell(row=r, column=4).value
    entry_dir = ws.cell(row=r, column=5).value
    vol = float(ws.cell(row=r, column=6).value or 0)
    price = float(ws.cell(row=r, column=7).value or 0)
    order = ws.cell(row=r, column=8).value
    profit = float(ws.cell(row=r, column=11).value or 0)
    comment = str(ws.cell(row=r, column=13).value or '')
    deals.append({
        'ticket': d_ticket,
        'order': order,
        'time': str(t_str),
        'type': typ,
        'dir': entry_dir,
        'vol': vol,
        'price': price,
        'profit': profit,
        'comment': comment
    })

print(f'Total orders: {len(orders)}, Total deals: {len(deals)}')

# Match in and out deals using FIFO queue for buy and sell
buy_open = []
sell_open = []
completed_trades = []

for d in deals:
    if d['dir'] == 'in':
        # Find matching order for SL
        sl = None
        for o in orders:
            if o['ticket'] == d['order']:
                sl = o['sl']
                break
        d['sl'] = sl
        if d['type'] == 'buy':
            buy_open.append(d)
        else:
            sell_open.append(d)
    elif d['dir'] == 'out':
        # Closing a position
        if d['type'] == 'sell': # closing a buy
            if buy_open:
                entry_deal = buy_open.pop(0)
                completed_trades.append({
                    'type': 'BUY',
                    'in_time': entry_deal['time'],
                    'out_time': d['time'],
                    'in_price': entry_deal['price'],
                    'out_price': d['price'],
                    'sl': entry_deal['sl'],
                    'profit': d['profit'],
                    'comment': d['comment']
                })
        elif d['type'] == 'buy': # closing a sell
            if sell_open:
                entry_deal = sell_open.pop(0)
                completed_trades.append({
                    'type': 'SELL',
                    'in_time': entry_deal['time'],
                    'out_time': d['time'],
                    'in_price': entry_deal['price'],
                    'out_price': d['price'],
                    'sl': entry_deal['sl'],
                    'profit': d['profit'],
                    'comment': d['comment']
                })

print(f'Total completed trades: {len(completed_trades)}')

# Analyze trade statistics
profits = [t['profit'] for t in completed_trades]
wins = [p for p in profits if p > 0]
losses = [p for p in profits if p <= 0]

print(f'Win count: {len(wins)} ({len(wins)/len(profits)*100:.1f}%), Loss count: {len(losses)} ({len(losses)/len(profits)*100:.1f}%)')
print(f'Avg Win: ${sum(wins)/len(wins):.2f}, Avg Loss: ${sum(losses)/len(losses):.2f}')
print(f'Total Profit: ${sum(profits):.2f}')

# SL Distances analysis
sl_dists = []
for t in completed_trades:
    if t['sl']:
        if t['type'] == 'BUY':
            dist = (t['in_price'] - t['sl'])
        else:
            dist = (t['sl'] - t['in_price'])
        sl_dists.append(dist)

if sl_dists:
    print(f'SL distances (USD on Gold): min=${min(sl_dists):.2f}, max=${max(sl_dists):.2f}, avg=${sum(sl_dists)/len(sl_dists):.2f}')

# Print 15 sample trades
print('\n=== Sample Trades ===')
for i, t in enumerate(completed_trades[10:25]):
    print('%2d: %-4s | In: %s @ %7.2f (SL: %s) -> Out: %s @ %7.2f | PnL: $%6.2f | %s' % 
          (i+1, t['type'], t['in_time'], t['in_price'], str(t['sl']), t['out_time'], t['out_price'], t['profit'], t['comment']))
