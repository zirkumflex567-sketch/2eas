import openpyxl
import pandas as pd
import numpy as np
from datetime import datetime

file_path = r'C:\Users\zirku\Downloads\ReportTester-53087839.xlsx'
wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb['Sheet1']

# Parse Deals / Trades starting at row 6600
trades_data = []
for r in range(6600, ws.max_row + 1):
    time_val = ws.cell(r, 1).value
    if not time_val:
        break
    trade_id = ws.cell(r, 2).value
    symbol = ws.cell(r, 3).value
    order_type = ws.cell(r, 4).value
    direction = ws.cell(r, 5).value # in, out, in/out
    volume = ws.cell(r, 6).value
    price = ws.cell(r, 7).value
    order_id = ws.cell(r, 8).value
    commission = ws.cell(r, 9).value or 0.0
    swap = ws.cell(r, 10).value or 0.0
    profit = ws.cell(r, 11).value or 0.0
    balance = ws.cell(r, 12).value
    comment = ws.cell(r, 13).value
    
    trades_data.append({
        'time': str(time_val),
        'trade_id': trade_id,
        'symbol': symbol,
        'type': order_type,
        'direction': str(direction),
        'volume': float(volume or 0),
        'price': float(price or 0),
        'commission': float(commission or 0),
        'swap': float(swap or 0),
        'profit': float(profit or 0),
        'balance': float(balance or 0) if balance is not None else None,
        'comment': str(comment or '')
    })

df = pd.DataFrame(trades_data)
print(f"Total deals parsed: {len(df)}")
df['datetime'] = pd.to_datetime(df['time'], format='%Y.%m.%d %H:%M:%S', errors='coerce')
df['date'] = df['datetime'].dt.date

# Filter closing deals (direction 'out')
df_out = df[df['direction'].str.contains('out', case=False, na=False)].copy()
df_out['net_pnl'] = df_out['profit'] + df_out['commission'] + df_out['swap']
print(f"Total closed deals: {len(df_out)}")

# Group by day to see daily balance changes
daily_pnl = df_out.groupby('date')['net_pnl'].sum()
print("\nTop 10 Worst Days by Closed PnL:")
print(daily_pnl.sort_values().head(10))

# Daily drawdown simulation:
# On each day, starting balance vs lowest closed balance
daily_stats = []
current_balance = 10000.0 # initial deposit
peak_balance = 10000.0

# Group all deals by date
for date, day_deals in df.groupby('date'):
    day_start_balance = current_balance
    day_min_balance = current_balance
    day_pnl = 0.0
    
    for _, row in day_deals.iterrows():
        if row['balance'] is not None and row['balance'] > 0:
            current_balance = row['balance']
            if current_balance < day_min_balance:
                day_min_balance = current_balance
            if current_balance > peak_balance:
                peak_balance = current_balance
        elif row['direction'] == 'out':
            current_balance += row['net_pnl']
            if current_balance < day_min_balance:
                day_min_balance = current_balance
            if current_balance > peak_balance:
                peak_balance = current_balance
                
    day_loss = day_start_balance - day_min_balance
    day_dd_pct = (day_loss / day_start_balance) * 100.0 if day_start_balance > 0 else 0
    # Also relative to initial balance (10,000):
    day_dd_initial_pct = (day_loss / 10000.0) * 100.0
    
    daily_stats.append({
        'date': date,
        'start_balance': day_start_balance,
        'end_balance': current_balance,
        'min_balance': day_min_balance,
        'day_loss': day_loss,
        'day_dd_pct': day_dd_pct,
        'day_dd_initial_pct': day_dd_initial_pct,
        'day_pnl': current_balance - day_start_balance
    })

daily_df = pd.DataFrame(daily_stats)
worst_daily = daily_df.sort_values(by='day_dd_pct', ascending=False).head(15)
print("\n--- Worst 15 Daily Drawdowns (Closed Deals) ---")
print(worst_daily[['date', 'start_balance', 'end_balance', 'day_loss', 'day_dd_pct', 'day_dd_initial_pct']].to_string())

# Count days exceeding limits
days_over_3_5 = len(daily_df[daily_df['day_dd_initial_pct'] >= 3.5])
days_over_4_0 = len(daily_df[daily_df['day_dd_initial_pct'] >= 4.0])
days_over_5_0 = len(daily_df[daily_df['day_dd_initial_pct'] >= 5.0])

print(f"\nDays violating Daily DD limit >= 3.5% (initial): {days_over_3_5}")
print(f"Days violating Daily DD limit >= 4.0% (initial): {days_over_4_0}")
print(f"Days violating Daily DD limit >= 5.0% (initial): {days_over_5_0}")

# Max peak-to-trough balance drawdown
daily_df['peak'] = daily_df['end_balance'].cummax()
daily_df['balance_dd'] = daily_df['peak'] - daily_df['end_balance']
daily_df['balance_dd_pct'] = (daily_df['balance_dd'] / daily_df['peak']) * 100.0
max_bal_dd = daily_df['balance_dd_pct'].max()
max_bal_dd_usd = daily_df['balance_dd'].max()
print(f"\nMax Balance Drawdown Peak-to-Trough: ${max_bal_dd_usd:.2f} ({max_bal_dd:.2f}%)")
