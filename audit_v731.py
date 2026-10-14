import openpyxl
import pandas as pd
import numpy as np
from datetime import datetime

file_path = r'C:/Users/zirku/.gemini/antigravity/brain/c98ddc14-d641-4466-a351-59d533463650/.user_uploaded/media_1791969819895_b1db4f88.xlsx'

wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb.active

# 1. Summary stats
orders_start_row = None
trades_start_row = None

for r in range(1, ws.max_row + 1):
    c1 = ws.cell(r, 1).value
    c1_str = str(c1).strip() if c1 is not None else ''
    if c1_str == 'Orders':
        orders_start_row = r
    elif c1_str == 'Trades':
        trades_start_row = r
        break

print(f"Orders row: {orders_start_row}, Trades row: {trades_start_row}")

# Parse Trades
header_row = trades_start_row + 1
trades = []

for r in range(header_row + 1, ws.max_row + 1):
    c1 = ws.cell(r, 1).value
    if c1 is None:
        continue
    c1_str = str(c1).strip()
    if not c1_str or c1_str.startswith('==='):
        continue
    
    row_vals = [ws.cell(r, c).value for c in range(1, 15)]
    time_val = str(row_vals[0]).strip()
    trade_id = row_vals[1]
    symbol = row_vals[2]
    order_type = row_vals[3]
    direction = str(row_vals[4] or '').strip().lower()
    volume = row_vals[5]
    price = row_vals[6]
    order_id = row_vals[7]
    commission = row_vals[8] or 0.0
    swap = row_vals[9] or 0.0
    profit = row_vals[10] or 0.0
    balance = row_vals[11]
    comment = str(row_vals[12] or '').strip()

    try:
        dt = datetime.strptime(time_val, '%Y.%m.%d %H:%M:%S')
    except Exception:
        continue

    trades.append({
        'time': time_val,
        'datetime': dt,
        'trade_id': trade_id,
        'symbol': symbol,
        'type': order_type,
        'direction': direction,
        'volume': float(volume or 0),
        'price': float(price or 0),
        'commission': float(commission or 0),
        'swap': float(swap or 0),
        'profit': float(profit or 0),
        'balance': float(balance) if balance is not None else None,
        'comment': comment
    })

df = pd.DataFrame(trades)
print(f"Total trade rows: {len(df)}")

deposit_row = df[(df['trade_id'] == 1) | (df['direction'] == '')]
initial_balance = deposit_row['balance'].iloc[0] if len(deposit_row) > 0 else 10000.0

df_out = df[df['direction'] == 'out'].copy()
df_out['net_pnl'] = df_out['profit'] + df_out['commission'] + df_out['swap']
df_out['date'] = df_out['datetime'].dt.date

print(f"Total Exits: {len(df_out)}")
print(f"Net Profit: ${df_out['net_pnl'].sum():,.2f}")
print(f"Final Balance: ${df_out['balance'].iloc[-1]:,.2f}")

# 2. Check 10% Max Drawdown
# 10% max DD on $10,000 initial is $1,000 max DD (Floor = $9,000)
# Or if based on trailing peak, check peak-to-trough
min_balance = df_out['balance'].min()
min_balance_row = df_out.loc[df_out['balance'].idxmin()]
lowest_equity = initial_balance - 342.76 # from MT5 "Rückgang Equity absolut: 342.76"

print("\n" + "=" * 80)
print("             1. AUDIT: MAX DRAWDOWN (10.0% GRENZE)")
print("=" * 80)
print(f"Ersteinlage:                                      ${initial_balance:,.2f}")
print(f"Statischer 10% Floor (Ersteinlage - 10%):         $9,000.00 (-$1,000.00)")
print(f"Statischer 9.100 $ Floor (-9%):                   $9,100.00 (-$900.00)")
print(f"Niedrigster Kontostand (Closed Balance):          ${min_balance:,.2f} (am {min_balance_row['time']})")
print(f"  -> Puffer zum 10% Floor ($9.000):               +${min_balance - 9000.0:,.2f} ÜBER DEM LIMIT")
print(f"  -> Puffer zum 9.100 $ Floor:                    +${min_balance - 9100.0:,.2f} ÜBER DEM LIMIT")
print(f"Niedrigste jemals erreichte Floating Equity:      ${lowest_equity:,.2f}")
print(f"  -> Puffer zum 10% Floor ($9.000):               +${lowest_equity - 9000.0:,.2f} ÜBER DEM LIMIT")
print(f"  -> Puffer zum 9.100 $ Floor:                    +${lowest_equity - 9100.0:,.2f} ÜBER DEM LIMIT")
print(f"Max. Rückgang Kontostand vom Peak (MT5 Report):   8.97% ($1,324.41)")
print(f"Max. Rückgang Equity vom Peak (MT5 Report):       11.67% ($1,766.54)")

# 3. Check Daily Drawdown (4.0%)
print("\n" + "=" * 80)
print("             2. AUDIT: DAILY DRAWDOWN (4.0% GRENZE)")
print("=" * 80)

daily_stats = []
running_balance = initial_balance

for date, group in df_out.groupby('date'):
    start_bal = running_balance
    end_bal = group['balance'].iloc[-1]
    running_balance = end_bal
    
    min_bal_day = group['balance'].min()
    loss_from_start = start_bal - min_bal_day
    loss_pct_start = (loss_from_start / start_bal) * 100.0 if start_bal > 0 else 0.0
    loss_pct_initial = (loss_from_start / initial_balance) * 100.0
    day_pnl = group['net_pnl'].sum()
    
    daily_stats.append({
        'date': date,
        'start_balance': start_bal,
        'end_balance': end_bal,
        'min_balance': min_bal_day,
        'day_pnl': day_pnl,
        'max_loss_from_start_usd': max(0.0, loss_from_start),
        'max_loss_pct_start': max(0.0, loss_pct_start),
        'max_loss_pct_initial': max(0.0, loss_pct_initial),
        'trades': len(group)
    })

df_days = pd.DataFrame(daily_stats)
df_worst = df_days.sort_values(by='max_loss_from_start_usd', ascending=False)

print(f"{'Datum':<12} | {'Start Balance':<14} | {'Min Balance':<14} | {'Tagesverlust':<14} | {'DD (% Start)':<12} | {'DD (% Init)':<12} | {'Status (BB)':<12} | {'Status (Fix 400$)':<18}")
print("-" * 115)

for _, r in df_worst.head(15).iterrows():
    status_bb = "OK" if r['max_loss_pct_start'] < 4.0 else "BREACH"
    status_fix = "OK" if r['max_loss_from_start_usd'] <= 400.0 else f"BREACH (+${r['max_loss_from_start_usd']-400.0:.2f})"
    print(f"{str(r['date']):<12} | ${r['start_balance']:>12,.2f} | ${r['min_balance']:>12,.2f} | ${r['max_loss_from_start_usd']:>12,.2f} | {r['max_loss_pct_start']:>10.2f}% | {r['max_loss_pct_initial']:>10.2f}% | {status_bb:<12} | {status_fix:<18}")

breaches_bb = df_days[df_days['max_loss_pct_start'] >= 4.0]
breaches_fix = df_days[df_days['max_loss_from_start_usd'] > 400.0]

print("\n" + "-" * 80)
print(f"Tage mit Tagesverlust >= 4.0% vom Tagesstart (Blueberry Funded Regel): {len(breaches_bb)}")
print(f"Tage mit Tagesverlust > 400.00 USD (Fixe 4% von Ersteinlage):         {len(breaches_fix)}")

max_day_dd_start = df_days['max_loss_pct_start'].max()
worst_day_start = df_days.loc[df_days['max_loss_pct_start'].idxmax()]
max_day_dd_usd = df_days['max_loss_from_start_usd'].max()
worst_day_usd = df_days.loc[df_days['max_loss_from_start_usd'].idxmax()]

print(f"\nHöchster relativer Tagesverlust: {max_day_dd_start:.2f}% (${worst_day_start['max_loss_from_start_usd']:,.2f}) am {worst_day_start['date']}")
print(f"Höchster absoluter Tagesverlust: ${max_day_dd_usd:,.2f} ({worst_day_usd['max_loss_pct_start']:.2f}% vom Tag) am {worst_day_usd['date']}")
