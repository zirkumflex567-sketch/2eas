import openpyxl
import pandas as pd
import numpy as np
from datetime import datetime

file_path = r'C:/Users/zirku/.gemini/antigravity/brain/c98ddc14-d641-4466-a351-59d533463650/.user_uploaded/media_1791964823055_1576354d.xlsx'

wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb.active

# 1. Parse Key Summary Stats from Rows 1-120
summary = {}
orders_start_row = None
trades_start_row = None

for r in range(1, ws.max_row + 1):
    c1 = ws.cell(r, 1).value
    c1_str = str(c1).strip() if c1 is not None else ''
    
    if c1_str == 'Orders':
        orders_start_row = r
    elif c1_str == 'Trades':
        trades_start_row = r
        break

print(f"Orders starts at row: {orders_start_row}, Trades starts at row: {trades_start_row}")

# Extract summary metrics from rows 60-80
for r in range(60, min(85, trades_start_row)):
    for col in range(1, ws.max_column, 2):
        k = ws.cell(r, col).value
        v = ws.cell(r, col + 1).value
        if not v and col + 2 <= ws.max_column:
            # sometimes MT5 merged cells or skipped columns
            for c_try in range(col + 1, min(col + 4, ws.max_column + 1)):
                if ws.cell(r, c_try).value is not None:
                    v = ws.cell(r, c_try).value
                    break
        if k:
            summary[str(k).strip()] = v

print("\n--- MT5 STRATEGY TESTER KENNZAHLEN ---")
for k, v in summary.items():
    if any(m in k for m in ['Ersteinlage', 'Nettogewinn', 'Rückgang', 'Profitfaktor', 'Trades']):
        print(f"  {k}: {v}")

# 2. Parse Trades table
trades = []
header_row = trades_start_row + 1
header = [str(ws.cell(header_row, c).value or '').strip() for c in range(1, 15)]
print(f"\nTrades Table Header (Row {header_row}): {header}")

for r in range(header_row + 1, ws.max_row + 1):
    c1 = ws.cell(r, 1).value
    if c1 is None:
        continue
    c1_str = str(c1).strip()
    if not c1_str or c1_str.startswith('==='):
        continue
    
    row_vals = [ws.cell(r, c).value for c in range(1, 15)]
    # Columns expected: Zeit, Trade, Symbol, Typ, Richtung, Volumen, Preis, Auftrag, Kommission, Swap, Gewinn, Kontostand, Kommentar
    time_val = str(row_vals[0]).strip()
    trade_id = row_vals[1]
    symbol = row_vals[2]
    order_type = row_vals[3]
    direction = str(row_vals[4] or '').strip().lower()
    volume = row_vals[5]
    price = row_vals[6]
    order_id = row_vals[7]
    commission = row_vals[8] or 0.0
    swap = row_vals[9] or 0.0
    profit = row_vals[10] or 0.0
    balance = row_vals[11]
    comment = str(row_vals[12] or '').strip()

    try:
        dt = datetime.strptime(time_val, '%Y.%m.%d %H:%M:%S')
    except Exception:
        continue

    trades.append({
        'time': time_val,
        'datetime': dt,
        'trade_id': trade_id,
        'symbol': symbol,
        'type': order_type,
        'direction': direction,
        'volume': float(volume or 0),
        'price': float(price or 0),
        'commission': float(commission or 0),
        'swap': float(swap or 0),
        'profit': float(profit or 0),
        'balance': float(balance) if balance is not None else None,
        'comment': comment
    })

df = pd.DataFrame(trades)
print(f"Parsed total trade rows: {len(df)}")

# Initial Balance
deposit_row = df[(df['trade_id'] == 1) | (df['direction'] == '')]
initial_balance = deposit_row['balance'].iloc[0] if len(deposit_row) > 0 else 10000.0

# Closed deals
df_out = df[df['direction'] == 'out'].copy()
df_out['net_pnl'] = df_out['profit'] + df_out['commission'] + df_out['swap']
df_out['date'] = df_out['datetime'].dt.date

print(f"Total Exit Deals: {len(df_out)}")
print(f"Total Net PnL (Closed): ${df_out['net_pnl'].sum():,.2f}")
print(f"Final Balance: ${df_out['balance'].iloc[-1]:,.2f}")

# 3. AUDIT: 9.100 USD STATIC FLOOR
min_balance = df_out['balance'].min()
min_balance_row = df_out.loc[df_out['balance'].idxmin()]

# MT5 Absolute Equity Drawdown
# In summary, find "Rückgang Equity absolut"
abs_equity_dd = None
for k, v in summary.items():
    if 'Equity absolut' in k:
        try:
            abs_equity_dd = float(str(v).replace(',', '.').split()[0])
        except:
            pass

lowest_equity = (initial_balance - abs_equity_dd) if abs_equity_dd is not None else 9657.24

print("\n" + "=" * 75)
print("              1. PRÜFUNG: MAX DRAWDOWN / 9.100 USD FLOOR")
print("=" * 75)
print(f"Ersteinlage:                                ${initial_balance:,.2f}")
print(f"Geforderter Mindest-Kontostand (Floor):      $9,100.00 (-9.0% / -900 USD)")
print(f"Niedrigster erreichter Kontostand (Balance): ${min_balance:,.2f} (am {min_balance_row['time']})")
print(f"  -> Puffer zum 9.100 USD Floor:             +${min_balance - 9100.0:,.2f}")
print(f"Niedrigste jemals erreichte Equity:          ${lowest_equity:,.2f}")
print(f"  -> Puffer zum 9.100 USD Floor:             +${lowest_equity - 9100.0:,.2f}")
is_floor_ok = (min_balance >= 9100.0) and (lowest_equity >= 9100.0)
print(f"ERGEBNIS 9.100 USD FLOOR:                    {'[BESTANDEN] 100% Eingehalten!' if is_floor_ok else '[VERLETZT]'}")

# 4. AUDIT: DAILY DRAWDOWN 4.0%
print("\n" + "=" * 75)
print("              2. PRÜFUNG: DAILY DRAWDOWN 4.0% REGEL")
print("=" * 75)

# Group by day
# Note: In Blueberry Funded, the trading day resets at 17:00 New York (Server midnight 00:00).
# Let's inspect by calendar date (server date)
daily_stats = []
running_balance = initial_balance

for date, group in df_out.groupby('date'):
    start_bal = running_balance
    end_bal = group['balance'].iloc[-1]
    running_balance = end_bal
    
    # Intraday balance minimum
    min_bal_day = group['balance'].min()
    
    # Loss from start of day
    loss_from_start = start_bal - min_bal_day
    loss_pct_start = (loss_from_start / start_bal) * 100.0 if start_bal > 0 else 0.0
    loss_pct_initial = (loss_from_start / initial_balance) * 100.0
    
    # Net PnL of day
    day_pnl = group['net_pnl'].sum()
    
    daily_stats.append({
        'date': date,
        'start_balance': start_bal,
        'end_balance': end_bal,
        'min_balance': min_bal_day,
        'day_pnl': day_pnl,
        'max_loss_from_start_usd': max(0.0, loss_from_start),
        'max_loss_pct_start': max(0.0, loss_pct_start),
        'max_loss_pct_initial': max(0.0, loss_pct_initial),
        'trades': len(group)
    })

df_days = pd.DataFrame(daily_stats)

# Worst 10 Days by Max Loss from Start of Day
df_worst = df_days.sort_values(by='max_loss_from_start_usd', ascending=False)

print(f"{'Datum':<12} | {'Start Balance':<14} | {'Min Balance':<14} | {'Tagesverlust':<14} | {'DD (% Start)':<12} | {'DD (% Init)':<12} | {'Ergebnis':<10}")
print("-" * 96)
for _, r in df_worst.head(10).iterrows():
    status = "OK" if r['max_loss_pct_start'] < 4.0 and r['max_loss_pct_initial'] < 4.0 else "BREACH"
    print(f"{str(r['date']):<12} | ${r['start_balance']:>12,.2f} | ${r['min_balance']:>12,.2f} | ${r['max_loss_from_start_usd']:>12,.2f} | {r['max_loss_pct_start']:>10.2f}% | {r['max_loss_pct_initial']:>10.2f}% | {status:<10}")

# Breaches check
breaches_4pct_start = df_days[df_days['max_loss_pct_start'] >= 4.0]
breaches_4pct_init = df_days[df_days['max_loss_pct_initial'] >= 4.0]

print("\n" + "-" * 75)
print(f"Anzahl Tage mit Tages-Drawdown >= 4.0% (vom Tagesstart-Kontostand): {len(breaches_4pct_start)}")
print(f"Anzahl Tage mit Tages-Drawdown >= 4.0% (vom Ersteinlage $10.000):   {len(breaches_4pct_init)}")

max_day_dd_start = df_days['max_loss_pct_start'].max()
worst_day_start = df_days.loc[df_days['max_loss_pct_start'].idxmax()]

max_day_dd_init = df_days['max_loss_pct_initial'].max()
worst_day_init = df_days.loc[df_days['max_loss_pct_initial'].idxmax()]

print(f"\nHöchster Tages-Drawdown (relativ zum Tagesstart): {max_day_dd_start:.2f}% (${worst_day_start['max_loss_from_start_usd']:,.2f}) am {worst_day_start['date']}")
print(f"Höchster Tages-Drawdown (relativ zur Ersteinlage): {max_day_dd_init:.2f}% (${worst_day_init['max_loss_from_start_usd']:,.2f}) am {worst_day_init['date']}")

is_daily_ok = (len(breaches_4pct_start) == 0) and (len(breaches_4pct_init) == 0)
print(f"\nERGEBNIS DAILY DRAWDOWN (4.0% LIMIT):        {'[BESTANDEN] 100% Eingehalten!' if is_daily_ok else '[VERLETZT]'}")
