import openpyxl

p_titan = r'C:\Users\zirku\Downloads\mt5_strategy_research\rev\500er\500acc ohne schutz.xlsx'
p_market = r'C:\Users\zirku\Downloads\mt5_strategy_research\rev\500eurotest.xlsx'

wb_t = openpyxl.load_workbook(p_titan, data_only=True)
ws_t = wb_t.active

wb_m = openpyxl.load_workbook(p_market, data_only=True)
ws_m = wb_m.active

# Parse deals of Titan Pro
deals_t = []
for r in range(1517, ws_t.max_row+1):
    c1 = ws_t.cell(row=r, column=1).value
    if not c1 or not str(c1).startswith('2026'): continue
    direction = ws_t.cell(row=r, column=5).value
    profit = float(ws_t.cell(row=r, column=11).value or 0)
    comment = str(ws_t.cell(row=r, column=13).value or '')
    if direction == 'out':
        deals_t.append({'time': str(c1), 'profit': profit, 'comment': comment})

# Parse deals of Market EA
deals_m = []
for r in range(898, ws_m.max_row+1):
    c1 = ws_m.cell(row=r, column=1).value
    if not c1 or not str(c1).startswith('2026'): continue
    direction = ws_m.cell(row=r, column=5).value
    profit = float(ws_m.cell(row=r, column=11).value or 0)
    comment = str(ws_m.cell(row=r, column=13).value or '')
    if direction == 'out':
        deals_m.append({'time': str(c1), 'profit': profit, 'comment': comment})

print(f'Titan exits: {len(deals_t)}, Market exits: {len(deals_m)}')

# Find Titan Pro biggest single-loss events
worst_titan = sorted(deals_t, key=lambda x: x['profit'])[:20]
print('\n=== TITAN PRO 15 BIGGEST LOSSES ===')
for d in worst_titan[:15]:
    print('%s | Loss: $%7.2f | %s' % (d['time'], d['profit'], d['comment']))
