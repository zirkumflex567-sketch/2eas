import openpyxl

p_titan = r'C:\Users\zirku\Downloads\mt5_strategy_research\rev\500er\500acc ohne schutz.xlsx'
wb = openpyxl.load_workbook(p_titan, data_only=True)
ws = wb.active

trades_r = 1517
trades = []
for r in range(trades_r+1, ws.max_row+1):
    c1 = ws.cell(row=r, column=1).value
    if not c1 or not str(c1).startswith('2026'): continue
    direction = ws.cell(row=r, column=5).value
    profit = float(ws.cell(row=r, column=11).value or 0)
    comm = float(ws.cell(row=r, column=9).value or 0)
    swap = float(ws.cell(row=r, column=10).value or 0)
    comment = str(ws.cell(row=r, column=13).value or '')
    if direction == 'out':
        trades.append({
            'time': str(c1),
            'profit': profit + comm + swap,
            'comment': comment
        })

# Simulate capping SL at -$60 and -$50
print('=== COMPARISON ON 500 USD ACCOUNT ===')
print('1. Original Titan Pro:               Net: +$3,302.30 | PF: 1.19 | Max DD: $1,723')
print('2. Market EA (XAUUSD 5 minute):      Net: +$4,505.61 | PF: 1.78 | Max DD: $  654')

for cap in [-60.0, -50.0]:
    capped_pnl = [t['profit'] if t['profit'] > 0 else max(cap, t['profit']) for t in trades]
    net = sum(capped_pnl)
    wins = sum(p for p in capped_pnl if p > 0)
    losses = sum(p for p in capped_pnl if p <= 0)
    pf = abs(wins / losses)
    
    # Calculate drawdown
    bal = 500.0
    peak = 500.0
    max_dd = 0.0
    for p in capped_pnl:
        bal += p
        if bal > peak: peak = bal
        dd = peak - bal
        if dd > max_dd: max_dd = dd
        
    print('3. Titan with Max Loss Capped at $%2d: Net: +$%7.2f | PF: %4.2f | Max DD: $%5.2f' % 
          (abs(cap), net, pf, max_dd))
