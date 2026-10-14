import openpyxl

p_market = r'C:\Users\zirku\Downloads\mt5_strategy_research\rev\500eurotest.xlsx'
p_titan  = r'C:\Users\zirku\Downloads\mt5_strategy_research\rev\500er\500acc ohne schutz.xlsx'

def get_deal_entries(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    trades_r = None
    for r in range(1, ws.max_row+1):
        if ws.cell(row=r, column=1).value in ['Trades', 'Deals', 'Geschäfte']:
            trades_r = r
            break
    entries = []
    for r in range(trades_r+2, ws.max_row+1):
        c1 = ws.cell(row=r, column=1).value
        if not c1 or not str(c1).startswith('2026'):
            continue
        entry_dir = ws.cell(row=r, column=5).value
        if entry_dir == 'in':
            entries.append({
                'time': str(c1),
                'type': ws.cell(row=r, column=4).value,
                'price': float(ws.cell(row=r, column=7).value or 0),
                'order': ws.cell(row=r, column=8).value
            })
    return entries

e_m = get_deal_entries(p_market)
e_t = get_deal_entries(p_titan)

print(f'Total Market entries: {len(e_m)}')
print(f'Total Titan entries:  {len(e_t)}')

# Check how many timestamps match exactly
times_m = {e['time']: e for e in e_m}
times_t = {e['time']: e for e in e_t}

common_times = set(times_m.keys()).intersection(set(times_t.keys()))
print(f'Exactly matching entry timestamps: {len(common_times)}')

# Let's inspect some common vs market-only vs titan-only
print('\nFirst 20 Market Entries:')
for i, e in enumerate(e_m[:20]):
    matched = 'MATCHED' if e['time'] in times_t else 'MARKET_ONLY'
    print('%s | %-4s | %7.2f | %s' % (e['time'], e['type'], e['price'], matched))
